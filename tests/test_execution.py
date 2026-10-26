"""
Unit Tests for Phase 4 Execution, Safety Validator, Abstention Policy, and DOCX Export
"""

import os
import sys
import pytest
from docx import Document

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "src")))

from schema import UnifiedDocument, DocumentPage, DocumentElement, ElementStyle
from execution import (
    ActionSafetyValidator,
    SafeAbstentionPolicy,
    AbstentionDecision,
    DocumentStateExecutor
)
from export import DocxDocumentExporter


@pytest.fixture
def sample_document():
    elem1 = DocumentElement(
        id="heading_main",
        type="heading",
        page=0,
        bbox=[60.0, 50.0, 740.0, 90.0],
        text="Executive Overview",
        style=ElementStyle(font_size=16, bold=False, alignment="left")
    )
    elem2 = DocumentElement(
        id="para_1",
        type="paragraph",
        page=0,
        bbox=[60.0, 100.0, 740.0, 140.0],
        text="Financial earnings grew by 14.2% year-over-year.",
        style=ElementStyle(font_size=11, bold=False, alignment="left")
    )
    page = DocumentPage(page_id=0, width=800.0, height=1000.0, elements=[elem1, elem2])
    return UnifiedDocument(document_id="doc_test_exec", pages=[page])


def test_action_validator(sample_document):
    # Valid action
    is_valid, msg = ActionSafetyValidator.validate_action(
        sample_document, "heading_main", "SET_STYLE", "bold", True
    )
    assert is_valid is True

    # Invalid font size (> 72 pt)
    is_valid, msg = ActionSafetyValidator.validate_action(
        sample_document, "heading_main", "SET_STYLE", "font_size", 999
    )
    assert is_valid is False
    assert "out of allowable range" in msg

    # Non-existent element
    is_valid, msg = ActionSafetyValidator.validate_action(
        sample_document, "non_existent_elem", "SET_STYLE", "bold", True
    )
    assert is_valid is False


def test_abstention_policy():
    policy = SafeAbstentionPolicy(
        min_execute_confidence=0.65,
        min_execute_margin=0.15,
        min_confirm_confidence=0.35
    )

    # 1. High confidence & decisive margin -> EXECUTE
    res = policy.evaluate(grounding_probs=[0.85, 0.10, 0.05], op_prob=1.0, prop_prob=1.0, val_prob=1.0)
    assert res.decision == AbstentionDecision.EXECUTE
    assert res.margin == pytest.approx(0.75)

    # 2. Ambiguous margin (close second candidate) -> REQUEST_CONFIRMATION
    res = policy.evaluate(grounding_probs=[0.48, 0.45, 0.07], op_prob=1.0, prop_prob=1.0, val_prob=1.0)
    assert res.decision == AbstentionDecision.REQUEST_CONFIRMATION

    # 3. Very low confidence -> ABSTAIN
    res = policy.evaluate(grounding_probs=[0.20, 0.20, 0.20, 0.20, 0.20], op_prob=0.5, prop_prob=0.5, val_prob=0.5)
    assert res.decision == AbstentionDecision.ABSTAIN


def test_state_executor(sample_document):
    modified_doc, audit = DocumentStateExecutor.execute_action(
        sample_document,
        target_element_id="heading_main",
        operation="SET_STYLE",
        property_name="bold",
        value=True
    )

    assert audit["success"] is True
    assert audit["before"]["style"]["bold"] is False
    assert audit["after"]["style"]["bold"] is True

    # Check that original document was NOT mutated (deep-copy verification)
    _, orig_elem = sample_document.find_element("heading_main")
    assert orig_elem.style.bold is False

    # Check modified document was updated
    _, mod_elem = modified_doc.find_element("heading_main")
    assert mod_elem.style.bold is True


def test_docx_exporter(sample_document, tmp_path):
    out_file = str(tmp_path / "exported.docx")
    DocxDocumentExporter.export_to_file(sample_document, out_file)

    assert os.path.exists(out_file)
    assert os.path.getsize(out_file) > 1000

    # Verify that the generated file is a valid Word document by re-opening it
    reopened = Document(out_file)
    para_texts = [p.text for p in reopened.paragraphs]
    assert "Executive Overview" in para_texts
    assert "Financial earnings grew by 14.2% year-over-year." in para_texts


def test_compound_state_executor(sample_document, tmp_path):
    # Test compound updates with named color "yellow"
    modified_doc, audit = DocumentStateExecutor.execute_action(
        sample_document,
        target_element_id="heading_main",
        operation="SET_STYLE",
        updates={"font_size": 18, "bold": True, "color": "yellow"}
    )

    assert audit["success"] is True
    assert audit["after"]["style"]["font_size"] == 18
    assert audit["after"]["style"]["bold"] is True
    assert audit["after"]["style"]["color"] == "#EAB308"

    # Verify export to Word preserves yellow color mapping
    out_file = str(tmp_path / "compound_exported.docx")
    DocxDocumentExporter.export_to_file(modified_doc, out_file)
    assert os.path.exists(out_file)
    assert os.path.getsize(out_file) > 1000


def test_advanced_formatting_validation(sample_document):
    # Valid formatting
    is_valid, _ = ActionSafetyValidator.validate_action(
        sample_document, "para_1", "SET_STYLE",
        updates={
            "strikethrough": True,
            "indent": 30.0,
            "line_spacing": 2.0,
            "letter_spacing": 2.5,
            "font_family": "Times New Roman",
            "alignment": "center"
        }
    )
    assert is_valid is True

    # Invalid indent (> 300)
    is_valid, msg = ActionSafetyValidator.validate_action(
        sample_document, "para_1", "SET_STYLE", "indent", 500
    )
    assert is_valid is False
    assert "Indentation 500.0 out of allowable range" in msg

    # Invalid line spacing (< 0.5)
    is_valid, msg = ActionSafetyValidator.validate_action(
        sample_document, "para_1", "SET_STYLE", "line_spacing", 0.2
    )
    assert is_valid is False
    assert "Line spacing 0.2 out of allowable range" in msg

    # Invalid letter spacing (> 30.0)
    is_valid, msg = ActionSafetyValidator.validate_action(
        sample_document, "para_1", "SET_STYLE", "letter_spacing", 50.0
    )
    assert is_valid is False
    assert "Letter spacing 50.0 out of allowable range" in msg


def test_advanced_formatting_execution_and_docx_export(sample_document, tmp_path):
    modified_doc, audit = DocumentStateExecutor.execute_action(
        sample_document,
        target_element_id="para_1",
        operation="SET_STYLE",
        updates={
            "strikethrough": True,
            "underline": True,
            "indent": 36.0,
            "line_spacing": 2.0,
            "letter_spacing": 3.0,
            "font_family": "Arial",
            "alignment": "center",
            "font_size": 14
        }
    )

    assert audit["success"] is True
    style = audit["after"]["style"]
    assert style["strikethrough"] is True
    assert style["underline"] is True
    assert style["indent"] == 36.0
    assert style["line_spacing"] == 2.0
    assert style["letter_spacing"] == 3.0
    assert style["font_family"] == "Arial"
    assert style["alignment"] == "center"
    assert style["font_size"] == 14

    # Export to DOCX and inspect Word XML AST properties
    out_file = str(tmp_path / "advanced_exported.docx")
    DocxDocumentExporter.export_to_file(modified_doc, out_file)
    assert os.path.exists(out_file)

    reopened = Document(out_file)
    target_p = None
    for p in reopened.paragraphs:
        if "Financial earnings" in p.text:
            target_p = p
            break

    assert target_p is not None
    # Verify paragraph formatting
    assert target_p.paragraph_format.left_indent.pt == pytest.approx(36.0)
    assert target_p.paragraph_format.line_spacing == pytest.approx(2.0)
    # Verify run formatting
    target_run = target_p.runs[0]
    assert target_run.underline is True
    assert target_run.font.strike is True
    assert target_run.font.name == "Arial"
    assert target_run.font.size.pt == pytest.approx(14.0)


def test_vlm_intent_extraction_formatting():
    from models.vlm_engine import extract_intent_directives

    d1 = extract_intent_directives("Set the heading to 24pt bold and emerald green")
    assert d1["updates"]["font_size"] == 24
    assert d1["updates"]["bold"] is True
    assert d1["updates"]["color"] == "#059669"

    d2 = extract_intent_directives("Add strikethrough and underline to the heading")
    assert d2["updates"]["strikethrough"] is True
    assert d2["updates"]["underline"] is True

    d3 = extract_intent_directives("Indent paragraph by 30pt and double line spacing")
    assert d3["updates"]["indent"] == 30.0
    assert d3["updates"]["line_spacing"] == 2.0

    d4 = extract_intent_directives("Letter spacing 3pt with wide expanded tracking and center align")
    assert d4["updates"]["letter_spacing"] == 3.0
    assert d4["updates"]["alignment"] == "center"

    d5 = extract_intent_directives("Change font to Times New Roman and remove bold")
    assert d5["updates"]["font_family"] == "Times New Roman"
    assert d5["updates"]["bold"] is False


def test_native_office_rendering(sample_document):
    from render import render_native_office_page
    from PIL import Image

    img = render_native_office_page(sample_document, page_id=0, highlight_elem_id="heading_main")
    assert isinstance(img, Image.Image)
    assert img.width > 500
    assert img.height > 500


def test_undo_redo_checkpoint_stack(sample_document):
    from backend.session_store import (
        store_document,
        push_checkpoint,
        undo_checkpoint,
        redo_checkpoint,
        get_checkpoint_status
    )

    doc_id = "test_chk_doc"
    store_document(doc_id, sample_document)

    # Initial checkpoint status
    status = get_checkpoint_status(doc_id)
    assert status["undo_count"] == 0
    assert status["redo_count"] == 0

    # Save checkpoint before mutation 1
    push_checkpoint(doc_id, sample_document)
    mutated_1, _ = DocumentStateExecutor.execute_action(
        sample_document, "heading_main", "SET_STYLE", "bold", True
    )
    store_document(doc_id, mutated_1)

    status = get_checkpoint_status(doc_id)
    assert status["undo_count"] == 1
    assert status["redo_count"] == 0

    # Save checkpoint before mutation 2
    push_checkpoint(doc_id, mutated_1)
    mutated_2, _ = DocumentStateExecutor.execute_action(
        mutated_1, "heading_main", "SET_STYLE", "font_size", 28
    )
    store_document(doc_id, mutated_2)

    status = get_checkpoint_status(doc_id)
    assert status["undo_count"] == 2
    assert status["redo_count"] == 0

    # Test Undo (should revert to mutated_1)
    undone_1 = undo_checkpoint(doc_id)
    assert undone_1 is not None
    _, el = undone_1.find_element("heading_main")
    assert el.style.bold is True
    assert el.style.font_size == 16  # font size reverted from 28 to 16

    status = get_checkpoint_status(doc_id)
    assert status["undo_count"] == 1
    assert status["redo_count"] == 1

    # Test Undo again (should revert to original)
    undone_2 = undo_checkpoint(doc_id)
    assert undone_2 is not None
    _, el = undone_2.find_element("heading_main")
    assert el.style.bold is False

    status = get_checkpoint_status(doc_id)
    assert status["undo_count"] == 0
    assert status["redo_count"] == 2

    # Test Redo (should reapply mutation 1)
    redone_1 = redo_checkpoint(doc_id)
    assert redone_1 is not None
    _, el = redone_1.find_element("heading_main")
    assert el.style.bold is True

    status = get_checkpoint_status(doc_id)
    assert status["undo_count"] == 1
    assert status["redo_count"] == 1


def test_target_text_substring_span_and_case(sample_document):
    # Element text: "Executive Overview"
    # Target only "Overview" and change color to blue
    mod_doc, audit = DocumentStateExecutor.execute_action(
        sample_document,
        target_element_id="heading_main",
        operation="SET_STYLE",
        property_name="color",
        value="#2563EB",
        target_text="Overview"
    )

    assert audit["success"] is True
    _, mod_el = mod_doc.find_element("heading_main")
    # Base element color should NOT be mutated
    assert mod_el.style.color != "#2563EB"
    # Metadata should contain the styled span
    assert "spans" in mod_el.metadata
    assert len(mod_el.metadata["spans"]) == 1
    assert mod_el.metadata["spans"][0]["text"] == "Overview"
    assert mod_el.metadata["spans"][0]["style"]["color"] == "#2563EB"

    # Test text case transformation on substring "Executive" -> uppercase
    case_doc, case_audit = DocumentStateExecutor.execute_action(
        mod_doc,
        target_element_id="heading_main",
        operation="SET_STYLE",
        property_name="text_case",
        value="uppercase",
        target_text="Executive"
    )
    assert case_audit["success"] is True
    _, case_el = case_doc.find_element("heading_main")
    assert "EXECUTIVE Overview" == case_el.text


def test_subscript_superscript_and_clear_style(sample_document):
    # 1. Test Subscript
    sub_doc, sub_audit = DocumentStateExecutor.execute_action(
        sample_document,
        target_element_id="para_1",
        operation="SET_STYLE",
        property_name="subscript",
        value=True
    )
    assert sub_audit["success"] is True
    _, sub_el = sub_doc.find_element("para_1")
    assert sub_el.style.subscript is True
    assert sub_el.style.superscript is False

    # 2. Test Superscript (should deactivate subscript)
    sup_doc, sup_audit = DocumentStateExecutor.execute_action(
        sub_doc,
        target_element_id="para_1",
        operation="SET_STYLE",
        property_name="superscript",
        value=True
    )
    assert sup_audit["success"] is True
    _, sup_el = sup_doc.find_element("para_1")
    assert sup_el.style.superscript is True
    assert sup_el.style.subscript is False

    # 3. Test Clear Style
    clear_doc, clear_audit = DocumentStateExecutor.execute_action(
        sup_doc,
        target_element_id="para_1",
        operation="CLEAR_STYLE"
    )
    assert clear_audit["success"] is True
    _, clear_el = clear_doc.find_element("para_1")
    assert clear_el.style.superscript is False
    assert clear_el.style.subscript is False
    assert "spans" not in clear_el.metadata


def test_canvas_multi_span_rendering(sample_document):
    from render.canvas import DocumentCanvasRenderer
    # Add a styled substring span to heading_main
    mod_doc, _ = DocumentStateExecutor.execute_action(
        sample_document,
        target_element_id="heading_main",
        operation="SET_STYLE",
        property_name="color",
        value="#2563EB",
        target_text="Overview"
    )
    renderer = DocumentCanvasRenderer()
    img = renderer.render_page(mod_doc, page_id=0)
    assert img is not None
    assert img.size == (800, 1000)




