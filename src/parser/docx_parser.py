"""
High-Fidelity Word DOCX Document Parser Engine
Implements Slide 10: Stage 1 (Document Parser for Word .docx files).

Extracts structural elements (headings, paragraphs, tables) with true 2D physical layout
bounding boxes, typography, and crisp visual page rasterization (150 DPI) into the UnifiedDocument schema.
"""

import os
import re
import shutil
import tempfile
import subprocess
import logging
from typing import List, Dict, Any, Optional
import pymupdf
import mammoth
import weasyprint
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from schema import UnifiedDocument, DocumentPage, DocumentElement, ElementStyle
from .base import BaseDocumentParser

logger = logging.getLogger("DocxParser")

ALIGN_LOOKUP = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify"
}


MAMMOTH_STYLE_MAP = "\n".join([
    "p[style-name='Title'] => h1:fresh",
    "p[style-name='Subtitle'] => h2:fresh",
    "p[style-name='Heading 1'] => h1:fresh",
    "p[style-name='Heading 2'] => h2:fresh",
    "p[style-name='Heading 3'] => h3:fresh",
    "p[style-name='Heading 4'] => h4:fresh",
    "p[style-name='Heading 5'] => h5:fresh",
    "p[style-name='Heading 6'] => h6:fresh"
])


class DocxDocumentParser(BaseDocumentParser):
    """
    Parses native Microsoft Word (.docx) documents into the UnifiedDocument model
    using high-fidelity native office layout rendering, table extraction, and true coordinate extraction.
    """

    def __init__(self, render_dpi: int = 150):
        self.render_dpi = render_dpi

    def parse(self, file_path: str, output_image_dir: Optional[str] = None) -> UnifiedDocument:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Document file not found: {file_path}")

        doc_id = os.path.splitext(os.path.basename(file_path))[0]

        if output_image_dir:
            os.makedirs(output_image_dir, exist_ok=True)

        # 1. Primary Engine: Native LibreOffice Headless (1:1 Multi-Column, Tables, and Typography)
        soffice_bin = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice_bin:
            try:
                return self._parse_with_soffice(file_path, doc_id, output_image_dir, soffice_bin)
            except Exception as e:
                logger.warning(f"Native LibreOffice conversion failed ({e}), falling back to WeasyPrint layout engine.")

        # 2. Secondary Engine: WeasyPrint CSS3 Layout Pipeline
        try:
            return self._parse_with_layout_engine(file_path, doc_id, output_image_dir)
        except Exception as e:
            logger.warning(f"High-fidelity rendering failed ({e}), falling back to direct docx parser.")
            return self._fallback_direct_parse(file_path, doc_id, output_image_dir)

    def _parse_with_soffice(
        self,
        file_path: str,
        doc_id: str,
        output_image_dir: Optional[str],
        soffice_bin: str
    ) -> UnifiedDocument:
        """
        Converts DOCX to high-fidelity vector PDF via headless LibreOffice,
        then extracts exact 2D bounding boxes, multi-column layouts, tables, and 150 DPI page rasters via PyMuPDF.
        """
        with tempfile.TemporaryDirectory(prefix="soffice_pdf_") as temp_dir:
            cmd = [
                soffice_bin,
                "--headless",
                "--convert-to", "pdf",
                file_path,
                "--outdir", temp_dir
            ]
            res = subprocess.run(cmd, capture_output=True, text=True, timeout=75)
            if res.returncode != 0:
                raise RuntimeError(f"LibreOffice conversion exited with code {res.returncode}: {res.stderr}")

            # Locate the generated PDF
            base_name = os.path.splitext(os.path.basename(file_path))[0]
            expected_pdf = os.path.join(temp_dir, f"{base_name}.pdf")
            if not os.path.exists(expected_pdf):
                pdf_files = [f for f in os.listdir(temp_dir) if f.lower().endswith(".pdf")]
                if not pdf_files:
                    raise FileNotFoundError(f"LibreOffice did not produce a PDF in {temp_dir}")
                expected_pdf = os.path.join(temp_dir, pdf_files[0])

            # Extract underlying python-docx table rows to match borderless layout tables
            docx_tables_rows = []
            try:
                docx_obj = Document(file_path)
                for t in docx_obj.tables:
                    for r in t.rows:
                        cells = [c.text.strip() for c in r.cells if c.text.strip()]
                        if len(cells) >= 2:
                            docx_tables_rows.append(cells)
            except Exception:
                docx_tables_rows = []

            pdf_doc = pymupdf.open(expected_pdf)
            pages: List[DocumentPage] = []

            for p_idx in range(len(pdf_doc)):
                pdf_page = pdf_doc[p_idx]
                rect = pdf_page.rect
                page_w = float(rect.width)
                page_h = float(rect.height)

                # 1. Rasterize crisp 150 DPI visual page bitmap
                img_path = None
                if output_image_dir:
                    pix = pdf_page.get_pixmap(dpi=self.render_dpi)
                    img_filename = f"{doc_id}_page_{p_idx}.png"
                    img_path = os.path.join(output_image_dir, img_filename)
                    pix.save(img_path)

                raw_elements: List[DocumentElement] = []

                # 2. Extract structured tables
                table_finder = pdf_page.find_tables()
                table_rects = []
                for t_idx, t in enumerate(table_finder.tables):
                    table_rects.append(pymupdf.Rect(t.bbox))
                    rows_data = t.extract()
                    for r_idx, r in enumerate(t.rows):
                        row_cells = rows_data[r_idx] if r_idx < len(rows_data) else []
                        row_text = " | ".join(str(c or "").strip() for c in row_cells if str(c or "").strip())
                        if not row_text:
                            continue
                        row_bbox = [float(c) for c in r.bbox]
                        raw_elements.append(DocumentElement(
                            id=f"table_row_{p_idx}_{t_idx}_{r_idx}",
                            type="table_row",
                            page=p_idx,
                            bbox=row_bbox,
                            text=row_text,
                            style=ElementStyle(
                                font_size=10,
                                font_family="DejaVu Sans",
                                bold=(r_idx == 0),
                                alignment="left"
                            )
                        ))

                # 3. Extract text blocks and embedded images outside of tables
                text_page = pdf_page.get_text("dict", flags=pymupdf.TEXT_DEHYPHENATE)
                blocks = text_page.get("blocks", [])

                font_sizes = []
                for b in blocks:
                    if b.get("type") == 0:
                        for line in b.get("lines", []):
                            for span in line.get("spans", []):
                                font_sizes.append(span.get("size", 11.0))
                median_size = (sorted(font_sizes)[len(font_sizes) // 2]) if font_sizes else 11.0

                text_counter = 0
                for b_idx, b in enumerate(blocks):
                    b_type = b.get("type", 0)
                    b_bbox = [float(c) for c in b.get("bbox", [0, 0, 0, 0])]
                    b_rect = pymupdf.Rect(b_bbox)

                    if any(b_rect.intersects(t_rect) for t_rect in table_rects):
                        continue

                    if b_type == 1:
                        elem = DocumentElement(
                            id=f"img_{p_idx}_{b_idx}",
                            type="image",
                            page=p_idx,
                            bbox=b_bbox,
                            text="[Embedded Image / Figure]",
                            style=ElementStyle()
                        )
                        raw_elements.append(elem)
                        continue

                    block_lines = b.get("lines", [])
                    if not block_lines:
                        continue

                    block_text_parts = []
                    max_span_size = 0.0
                    is_bold_block = False
                    is_italic_block = False
                    font_family = "DejaVu Sans"

                    for line in block_lines:
                        line_spans = line.get("spans", [])
                        line_text = "".join(span.get("text", "") for span in line_spans).strip()
                        if line_text:
                            block_text_parts.append(line_text)

                        for span in line_spans:
                            span_size = span.get("size", 11.0)
                            if span_size > max_span_size:
                                max_span_size = span_size
                            flags = span.get("flags", 0)
                            if flags & (1 << 4):
                                is_bold_block = True
                            if flags & (1 << 1):
                                is_italic_block = True
                            if span.get("font"):
                                font_family = span.get("font")

                    full_text = " ".join(block_text_parts).strip()
                    if not full_text:
                        continue

                    # Check if text block matches any docx table row (for borderless tables)
                    matched_row = None
                    clean_lower = full_text.lower()
                    for r_cells in docx_tables_rows:
                        if all(c.lower() in clean_lower for c in r_cells):
                            matched_row = r_cells
                            break

                    if matched_row:
                        elem_type = "table_row"
                        full_text = " | ".join(matched_row)
                    elif max_span_size >= 13.5 or (is_bold_block and len(full_text) < 120 and max_span_size >= median_size):
                        elem_type = "heading"
                    else:
                        elem_type = "paragraph"

                    b_mid_x = (b_bbox[0] + b_bbox[2]) / 2.0
                    page_mid_x = page_w / 2.0
                    b_width = b_bbox[2] - b_bbox[0]
                    left_margin = b_bbox[0]
                    right_margin = page_w - b_bbox[2]
                    margin_diff = abs(left_margin - right_margin)
                    is_centered = (
                        abs(b_mid_x - page_mid_x) < (page_w * 0.06)
                        and (b_width < (page_w * 0.72) or margin_diff < (page_w * 0.08))
                        and left_margin > (page_w * 0.04)
                    )
                    if is_centered:
                        elem_alignment = "center"
                    else:
                        elem_alignment = "left"

                    elem = DocumentElement(
                        id=f"{elem_type}_{p_idx}_{text_counter}",
                        type=elem_type,
                        page=p_idx,
                        bbox=b_bbox,
                        text=full_text,
                        style=ElementStyle(
                            font_size=int(round(max_span_size)),
                            font_family=font_family,
                            bold=is_bold_block,
                            italic=is_italic_block,
                            alignment=elem_alignment
                        )
                    )
                    raw_elements.append(elem)
                    text_counter += 1

                # 4. Extract embedded images from PyMuPDF
                image_info_list = pdf_page.get_image_info(xrefs=True)
                image_elements: List[DocumentElement] = []
                for img_i, img_meta in enumerate(image_info_list):
                    img_bbox = [float(c) for c in img_meta.get("bbox", [0, 0, 0, 0])]
                    img_w = abs(img_bbox[2] - img_bbox[0])
                    img_h = abs(img_bbox[3] - img_bbox[1])
                    if img_w < 15.0 or img_h < 15.0:
                        continue  # Skip tiny icon or bullet artifacts
                    img_elem = DocumentElement(
                        id=f"image_{p_idx}_{img_i}",
                        type="image",
                        page=p_idx,
                        bbox=img_bbox,
                        text="[Embedded Figure / Architecture Diagram]",
                        style=ElementStyle()
                    )
                    image_elements.append(img_elem)

                # 5. Link figure captions with adjacent image diagrams
                for el in raw_elements:
                    if el.type in ("paragraph", "heading"):
                        cap_m = re.match(r"^\s*Fig(?:ure)?\.?\s*([0-9]+|[ivx]+|[a-z])", el.text, re.IGNORECASE)
                        if cap_m:
                            fig_n = cap_m.group(1).lower()
                            el.metadata["figure_number"] = fig_n
                            best_img_bbox = None
                            min_dist = 999.0
                            for im_el in image_elements:
                                ib = im_el.bbox
                                v_dist = min(abs(el.bbox[1] - ib[3]), abs(ib[1] - el.bbox[3]))
                                if v_dist < min_dist and v_dist <= 50.0:
                                    min_dist = v_dist
                                    best_img_bbox = ib
                            if best_img_bbox:
                                el.metadata["associated_image_bbox"] = best_img_bbox
                                el.metadata["composite_bbox"] = [
                                    round(min(el.bbox[0], best_img_bbox[0]), 2),
                                    round(min(el.bbox[1], best_img_bbox[1]), 2),
                                    round(max(el.bbox[2], best_img_bbox[2]), 2),
                                    round(max(el.bbox[3], best_img_bbox[3]), 2)
                                ]

                raw_elements.extend(image_elements)

                pages.append(DocumentPage(
                    page_id=p_idx,
                    width=page_w,
                    height=page_h,
                    elements=raw_elements,
                    image_path=img_path
                ))

            pdf_doc.close()
            return UnifiedDocument(document_id=doc_id, pages=pages)

    def _parse_with_layout_engine(self, file_path: str, doc_id: str, output_image_dir: Optional[str]) -> UnifiedDocument:
        """
        Converts DOCX to structured semantic HTML, renders to multi-page PDF via WeasyPrint,
        and extracts true 2D physical bounding boxes, table rows, headings, typography, and crisp visual page PNGs.
        """
        with open(file_path, "rb") as docx_file:
            conversion_res = mammoth.convert_to_html(docx_file, style_map=MAMMOTH_STYLE_MAP)
            body_html = conversion_res.value

        styled_html = f"""
        <!DOCTYPE html>
        <html>
        <head>
        <meta charset="utf-8">
        <style>
          @page {{
            size: A4;
            margin: 20mm;
          }}
          body {{
            font-family: "DejaVu Sans", "Noto Sans", -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif;
            font-size: 11pt;
            line-height: 1.5;
            color: #0F172A;
          }}
          h1, h2, h3, h4 {{
            color: #0F172A;
            font-weight: bold;
            margin-top: 1.2em;
            margin-bottom: 0.5em;
          }}
          h1 {{ font-size: 18pt; border-bottom: 2px solid #E2E8F0; padding-bottom: 4px; }}
          h2 {{ font-size: 14pt; }}
          h3 {{ font-size: 12pt; }}
          p {{ margin-bottom: 0.8em; }}
          table {{
            border-collapse: collapse;
            width: 100%;
            margin: 1em 0;
          }}
          table, th, td {{
            border: 1px solid #CBD5E1;
          }}
          th, td {{
            padding: 6px 10px;
            text-align: left;
            font-size: 10pt;
          }}
          th {{
            background-color: #F8FAFC;
            font-weight: bold;
          }}
        </style>
        </head>
        <body>{body_html}</body>
        </html>
        """

        pdf_bytes = weasyprint.HTML(string=styled_html).write_pdf()
        pdf_doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
        pages: List[DocumentPage] = []

        for p_idx in range(len(pdf_doc)):
            pdf_page = pdf_doc[p_idx]
            rect = pdf_page.rect
            page_w = float(rect.width)
            page_h = float(rect.height)

            # 1. Rasterize visual page bitmap (150 DPI)
            img_path = None
            if output_image_dir:
                pix = pdf_page.get_pixmap(dpi=self.render_dpi)
                img_filename = f"{doc_id}_page_{p_idx}.png"
                img_path = os.path.join(output_image_dir, img_filename)
                pix.save(img_path)

            raw_elements: List[DocumentElement] = []

            # 2. Extract structured tables as table_row elements
            table_finder = pdf_page.find_tables()
            table_rects = []
            for t_idx, t in enumerate(table_finder.tables):
                table_rects.append(pymupdf.Rect(t.bbox))
                rows_data = t.extract()
                for r_idx, r in enumerate(t.rows):
                    row_cells = rows_data[r_idx] if r_idx < len(rows_data) else []
                    row_text = " | ".join(str(c or "").strip() for c in row_cells if str(c or "").strip())
                    if not row_text:
                        continue
                    row_bbox = [float(c) for c in r.bbox]
                    raw_elements.append(DocumentElement(
                        id=f"table_row_{p_idx}_{t_idx}_{r_idx}",
                        type="table_row",
                        page=p_idx,
                        bbox=row_bbox,
                        text=row_text,
                        style=ElementStyle(
                            font_size=10,
                            font_family="DejaVu Sans",
                            bold=(r_idx == 0),
                            alignment="left"
                        )
                    ))

            # 3. Extract text blocks and embedded images outside of tables
            text_page = pdf_page.get_text("dict", flags=pymupdf.TEXT_DEHYPHENATE)
            blocks = text_page.get("blocks", [])

            # Compute median font size for relative heading detection
            font_sizes = []
            for b in blocks:
                if b.get("type") == 0:
                    for line in b.get("lines", []):
                        for span in line.get("spans", []):
                            font_sizes.append(span.get("size", 11.0))
            median_size = (sorted(font_sizes)[len(font_sizes) // 2]) if font_sizes else 11.0

            text_counter = 0
            for b_idx, b in enumerate(blocks):
                b_type = b.get("type", 0)
                b_bbox = [float(c) for c in b.get("bbox", [0, 0, 0, 0])]
                b_rect = pymupdf.Rect(b_bbox)

                # Skip blocks that overlap any extracted table
                if any(b_rect.intersects(t_rect) for t_rect in table_rects):
                    continue

                # Type 1: Image
                if b_type == 1:
                    elem = DocumentElement(
                        id=f"img_{p_idx}_{b_idx}",
                        type="image",
                        page=p_idx,
                        bbox=b_bbox,
                        text="[Embedded Image / Figure]",
                        style=ElementStyle()
                    )
                    raw_elements.append(elem)
                    continue

                # Type 0: Text Block
                block_lines = b.get("lines", [])
                if not block_lines:
                    continue

                block_text_parts = []
                max_span_size = 0.0
                is_bold_block = False
                is_italic_block = False
                font_family = "DejaVu Sans"

                for line in block_lines:
                    line_spans = line.get("spans", [])
                    line_text = "".join(span.get("text", "") for span in line_spans).strip()
                    if line_text:
                        block_text_parts.append(line_text)

                    for span in line_spans:
                        span_size = span.get("size", 11.0)
                        if span_size > max_span_size:
                            max_span_size = span_size
                        flags = span.get("flags", 0)
                        if flags & (1 << 4):  # Bold flag
                            is_bold_block = True
                        if flags & (1 << 1):  # Italic flag
                            is_italic_block = True
                        if span.get("font"):
                            font_family = span.get("font")

                full_text = " ".join(block_text_parts).strip()
                if not full_text:
                    continue

                # Classify semantic element type
                if max_span_size >= 13.5 or (is_bold_block and len(full_text) < 120 and max_span_size >= median_size):
                    elem_type = "heading"
                else:
                    elem_type = "paragraph"

                elem = DocumentElement(
                    id=f"{elem_type}_{p_idx}_{text_counter}",
                    type=elem_type,
                    page=p_idx,
                    bbox=b_bbox,
                    text=full_text,
                    style=ElementStyle(
                        font_size=int(round(max_span_size)),
                        font_family=font_family,
                        bold=is_bold_block,
                        italic=is_italic_block,
                        alignment="left"
                    )
                )
                raw_elements.append(elem)
                text_counter += 1

            # 4. Sort all elements top-to-bottom by vertical coordinate (bbox[1])
            raw_elements.sort(key=lambda el: el.bbox[1])

            pages.append(DocumentPage(
                page_id=p_idx,
                width=page_w,
                height=page_h,
                elements=raw_elements,
                image_path=img_path
            ))

        return UnifiedDocument(document_id=doc_id, pages=pages)

    def _fallback_direct_parse(self, file_path: str, doc_id: str, output_image_dir: Optional[str]) -> UnifiedDocument:
        """
        Fallback direct python-docx parsing if headless layout rendering is unavailable.
        """
        doc_obj = Document(file_path)
        elements: List[DocumentElement] = []
        curr_y = 60.0
        elem_counter = 0

        for p in doc_obj.paragraphs:
            text = p.text.strip()
            if not text:
                curr_y += 14.0
                continue

            style_name = p.style.name.lower() if p.style else ""
            if "heading" in style_name or "title" in style_name:
                elem_type = "heading"
                font_size = 18 if "title" in style_name else 14
                is_bold = True
                elem_h = 32.0
            else:
                elem_type = "paragraph"
                font_size = 11
                is_bold = False
                elem_h = max(24.0, (len(text) / 75.0) * 18.0)

            is_italic = False
            for run in p.runs:
                if run.bold:
                    is_bold = True
                if run.italic:
                    is_italic = True
                if run.font.size:
                    font_size = int(run.font.size.pt)

            align_str = ALIGN_LOOKUP.get(p.alignment, "left")
            elem_id = f"{elem_type}_0_{elem_counter}"
            elem_counter += 1
            bbox = [60.0, curr_y, 740.0, curr_y + elem_h]

            elem = DocumentElement(
                id=elem_id,
                type=elem_type,
                page=0,
                bbox=bbox,
                text=text,
                style=ElementStyle(
                    font_size=font_size,
                    bold=is_bold,
                    italic=is_italic,
                    alignment=align_str
                )
            )
            elements.append(elem)
            curr_y += elem_h + 12.0

        page = DocumentPage(page_id=0, width=800.0, height=max(1000.0, curr_y + 60.0), elements=elements)
        return UnifiedDocument(document_id=doc_id, pages=[page])
