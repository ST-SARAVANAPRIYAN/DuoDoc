"""
Native Microsoft Word (.docx) Document Exporter
Implements Slide 10: Stage 7 (Document Generation).

Serializes modified UnifiedDocument AST models back into fully editable,
native Microsoft Word (.docx) documents preserving styles, alignments, and tables.
"""

import io
import re
from dataclasses import asdict
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import Optional, List, Dict, Any

from schema import UnifiedDocument, DocumentPage, DocumentElement


ALIGN_MAPPING = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY
}


COLOR_EXPORT_MAP = {
    "yellow": "#EAB308",
    "gold": "#D97706",
    "amber": "#D97706",
    "red": "#DC2626",
    "blue": "#2563EB",
    "green": "#059669",
    "emerald": "#059669",
    "purple": "#7C3AED",
    "orange": "#EA580C",
    "black": "#0F172A",
    "gray": "#64748B",
    "grey": "#64748B"
}


def hex_to_rgb(hex_str: str) -> Optional[RGBColor]:
    if not hex_str:
        return None
    hex_clean = COLOR_EXPORT_MAP.get(str(hex_str).lower().strip(), str(hex_str).strip())
    if not hex_clean.startswith("#") or len(hex_clean) != 7:
        return None
    try:
        r = int(hex_clean[1:3], 16)
        g = int(hex_clean[3:5], 16)
        b = int(hex_clean[5:7], 16)
        return RGBColor(r, g, b)
    except Exception:
        return None


class DocxDocumentExporter:
    """
    Exports a UnifiedDocument to a native .docx file or in-memory byte stream.
    """

    @staticmethod
    def _apply_run_style(run, style_dict_or_obj):
        if hasattr(style_dict_or_obj, "bold"):
            bold = style_dict_or_obj.bold
            italic = style_dict_or_obj.italic
            underline = getattr(style_dict_or_obj, "underline", False)
            strike = getattr(style_dict_or_obj, "strikethrough", False)
            font_size = getattr(style_dict_or_obj, "font_size", 11)
            font_family = getattr(style_dict_or_obj, "font_family", None)
            color = getattr(style_dict_or_obj, "color", None)
        else:
            bold = style_dict_or_obj.get("bold")
            italic = style_dict_or_obj.get("italic")
            underline = style_dict_or_obj.get("underline")
            strike = style_dict_or_obj.get("strikethrough") or style_dict_or_obj.get("strike")
            font_size = style_dict_or_obj.get("font_size")
            font_family = style_dict_or_obj.get("font_family")
            color = style_dict_or_obj.get("color")

        if bold is not None:
            run.bold = bool(bold)
        if italic is not None:
            run.italic = bool(italic)
        if underline is not None:
            run.underline = bool(underline)
        if strike is not None:
            run.font.strike = bool(strike)
        if font_size:
            try:
                run.font.size = Pt(float(font_size))
            except Exception:
                pass
        if font_family and font_family != "Inter":
            run.font.name = font_family
        if color:
            rgb = hex_to_rgb(color)
            if rgb:
                run.font.color.rgb = rgb

    @staticmethod
    def _add_runs_to_paragraph(p, text: str, base_style, spans: Optional[List[Dict[str, Any]]] = None):
        if not spans:
            run = p.add_run(text)
            DocxDocumentExporter._apply_run_style(run, base_style)
            return

        matches = []
        for s in spans:
            stext = s.get("text", "")
            if not stext:
                continue
            for m in re.finditer(re.escape(stext), text, re.IGNORECASE):
                matches.append((m.start(), m.end(), s.get("style", {})))

        if not matches:
            run = p.add_run(text)
            DocxDocumentExporter._apply_run_style(run, base_style)
            return

        matches.sort(key=lambda x: (x[0], -x[1]))
        cur_idx = 0
        for start, end, sstyle in matches:
            if start < cur_idx:
                continue
            if start > cur_idx:
                r_base = p.add_run(text[cur_idx:start])
                DocxDocumentExporter._apply_run_style(r_base, base_style)
            r_span = p.add_run(text[start:end])
            base_dict = asdict(base_style) if hasattr(base_style, "__dataclass_fields__") else dict(base_style)
            merged_style = {**base_dict, **sstyle}
            DocxDocumentExporter._apply_run_style(r_span, merged_style)
            cur_idx = end

        if cur_idx < len(text):
            r_tail = p.add_run(text[cur_idx:])
            DocxDocumentExporter._apply_run_style(r_tail, base_style)

    @staticmethod
    def export_to_stream(doc: UnifiedDocument) -> io.BytesIO:
        """
        Converts the UnifiedDocument into an in-memory BytesIO stream of a .docx file.
        """
        word_doc = Document()

        for page in doc.pages:
            for elem in page.elements:
                style = elem.style
                text = elem.text.strip()
                if not text:
                    continue

                align_enum = ALIGN_MAPPING.get(style.alignment.lower(), WD_ALIGN_PARAGRAPH.LEFT)
                spans = elem.metadata.get("spans") if elem.metadata else None

                # Headings
                if elem.type == "heading":
                    p = word_doc.add_paragraph()
                    p.alignment = align_enum
                    if getattr(style, "indent", 0.0) > 0:
                        p.paragraph_format.left_indent = Pt(style.indent)
                    if getattr(style, "line_spacing", 1.15) > 0:
                        p.paragraph_format.line_spacing = style.line_spacing

                    DocxDocumentExporter._add_runs_to_paragraph(p, text, style, spans)

                # Table rows (pipe-separated)
                elif elem.type == "table_row":
                    cells = [c.strip() for c in text.split("|")]
                    table = word_doc.add_table(rows=1, cols=len(cells))
                    table.style = "Table Grid"
                    hdr_cells = table.rows[0].cells
                    for c_idx, cell_text in enumerate(cells):
                        hdr_cells[c_idx].text = cell_text
                        for p in hdr_cells[c_idx].paragraphs:
                            for r in p.runs:
                                r.font.size = Pt(max(9, style.font_size))
                                r.bold = style.bold

                # Standard Paragraphs
                else:
                    p = word_doc.add_paragraph()
                    p.alignment = align_enum
                    if getattr(style, "indent", 0.0) > 0:
                        p.paragraph_format.left_indent = Pt(style.indent)
                    if getattr(style, "line_spacing", 1.15) > 0:
                        p.paragraph_format.line_spacing = style.line_spacing

                    DocxDocumentExporter._add_runs_to_paragraph(p, text, style, spans)

        stream = io.BytesIO()
        word_doc.save(stream)
        stream.seek(0)
        return stream

    @staticmethod
    def export_to_file(doc: UnifiedDocument, output_path: str):
        stream = DocxDocumentExporter.export_to_stream(doc)
        with open(output_path, "wb") as f:
            f.write(stream.getbuffer())


# Convenience standalone function
def export_docx(doc: UnifiedDocument, output_path: Optional[str] = None) -> io.BytesIO:
    if output_path:
        DocxDocumentExporter.export_to_file(doc, output_path)
    return DocxDocumentExporter.export_to_stream(doc)
