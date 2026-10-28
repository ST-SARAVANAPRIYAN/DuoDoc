"""
Unit Tests for Document Parsers
"""

import os
import sys
import pytest
from docx import Document

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "src")))

from parser import DocxDocumentParser


def test_docx_parser(tmp_path):
    # Create temporary docx
    test_docx_path = str(tmp_path / "test_doc.docx")
    doc = Document()
    doc.add_heading("Deep Learning System Proposal", level=0)
    p = doc.add_paragraph("This document outlines the project architecture.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Parser"
    table.cell(1, 1).text = "Operational"
    
    doc.save(test_docx_path)

    # Parse with DocxDocumentParser
    parser = DocxDocumentParser()
    unified_doc = parser.parse(test_docx_path)

    assert unified_doc.document_id == "test_doc"
    assert len(unified_doc.pages) == 1
    page = unified_doc.pages[0]

    # Must contain heading, paragraph, and table rows
    elem_types = [el.type for el in page.elements]
    assert "heading" in elem_types
    assert "paragraph" in elem_types
    assert "table_row" in elem_types

    # Coordinates must be valid
    valid, errors = unified_doc.validate_integrity()
    assert valid is True
    assert len(errors) == 0


def test_pdf_parser(tmp_path):
    from parser import PdfDocumentParser
    pdf_path = "docs/final.pdf"
    if os.path.exists(pdf_path):
        parser = PdfDocumentParser(render_dpi=100)
        img_out = str(tmp_path / "pdf_renders")
        unified_doc = parser.parse(pdf_path, output_image_dir=img_out)

        assert unified_doc.document_id == "final"
        assert len(unified_doc.pages) == 12
        assert len(unified_doc.pages[0].elements) > 0

        valid, errors = unified_doc.validate_integrity()
        assert valid is True
        assert len(errors) == 0

